"""Export for sign-off and single-case regeneration."""
import io

import pytest
from openpyxl import load_workbook

from helper import regenerate
from helper.claude import ClaudeError, ClaudeStatus
from helper.runner import Helper

from .test_flows import FakeFlowClaude, run_flow_helper
from .test_unit_cases import FakeClaude, case, run_helper, setup_project


def unit_project(project):
    lead, pid, helper, endpoint_ids = setup_project(project)
    lead.post(f'/api/projects/{pid}/unit/generate', json={'endpointIds': [endpoint_ids['createPet']]})
    run_helper(helper, FakeClaude(drop=0))
    cases = lead.get(f'/api/projects/{pid}/unit/endpoints/{endpoint_ids["createPet"]}').json()['cases']
    return lead, pid, helper, endpoint_ids, cases


def test_excel_and_csv_export(project):
    lead, pid, helper, endpoint_ids, cases = unit_project(project)
    lead.post(f'/api/projects/{pid}/flows/generate')
    run_flow_helper(helper, FakeFlowClaude())
    lead.post(f'/api/projects/{pid}/unit/cases/review', json={'caseIds': [cases[0]['id']], 'action': 'approve'})

    response = lead.get(f'/api/projects/{pid}/export?module=all&format=xlsx')
    assert response.status_code == 200 and 'spreadsheetml' in response.headers['content-type']
    assert 'orders-api-all-cases-' in response.headers['content-disposition']
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ['Unit cases', 'Flow cases']
    unit = list(book['Unit cases'].values)
    assert unit[0][:3] == ('Case ID', 'Endpoint', 'Title') and len(unit) == len(cases) + 1
    assert unit[1][1] == 'POST /pets' and unit[1][5] == 'Approved'
    flows = list(book['Flow cases'].values)
    assert len(flows) == 7 and any('1. POST /pets' in row[7] and 'saves petId=$.id' in row[7] for row in flows[1:])

    approved = lead.get(f'/api/projects/{pid}/export?module=unit&format=csv&status=approved')
    text = approved.content.decode('utf-8-sig')
    assert approved.status_code == 200 and text.count('\n') == 2 and cases[0]['code'] in text
    assert lead.get(f'/api/projects/{pid}/export?module=all&format=csv').status_code == 400
    assert any('exported' in a['summary'] for a in lead.get(f'/api/projects/{pid}/activity').json())


def test_csv_neutralises_formulas():
    from studio.routers.cases import _csv
    assert _csv(['a'], [['=HYPERLINK("x")'], ['-1'], ['ok']]).decode('utf-8-sig').splitlines()[1:] == ['"\'=HYPERLINK(""x"")"', "'-1", 'ok']


class Rewriter:
    def __init__(self):
        self.tasks = []

    def __call__(self, agent, task, schema, **options):
        self.tasks.append((agent, task))
        return {'cases': [{**case(task['onlyKeys'][0], 'Rewritten case'), 'needsConfirmation': True, 'confirmationNote': 'Check 422'}]}


def test_regenerate_unit_case_uses_instruction_and_review_comment(project):
    lead, pid, helper, _, cases = unit_project(project)
    target = cases[2]
    lead.post(f'/api/projects/{pid}/unit/cases/review', json={'caseIds': [target['id']], 'action': 'needs_changes', 'comment': 'Use a 51-char name'})
    response = lead.post(f'/api/projects/{pid}/cases/{target["id"]}/regenerate', json={'instruction': 'Also assert the error code'})
    assert response.status_code == 200
    assert lead.post(f'/api/projects/{pid}/cases/{target["id"]}/regenerate', json={}).status_code == 409
    assert target['id'] in lead.get(f'/api/projects/{pid}/unit/endpoints/{_endpoint(lead, pid)}').json()['regenerating']

    fake = Rewriter()
    status = ClaudeStatus(installed=True, loggedIn=True)
    handlers = {'case.regenerate': lambda job_input, report, submit=None: regenerate.generate(job_input, report, submit, agent=fake)}
    assert Helper(helper, log=lambda _: None, claude_status=lambda: status, handlers=handlers).run_once()

    agent, task = fake.tasks[0]
    issues = [f['issue'] for f in task['reviewFindings']]
    assert agent == 'unit-case-writer' and task['mode'] == 'revise'
    assert 'QA instruction: Also assert the error code' in issues and 'Reviewer comment: Use a 51-char name' in issues
    updated = lead.get(f'/api/projects/{pid}/unit/cases/{target["id"]}').json()
    assert updated['title'] == 'Rewritten case' and updated['status'] == 'in_review' and updated['reviewComment'] == ''
    assert updated['needsConfirmation'] and updated['content']['confirmationNote'] == 'Check 422'
    assert updated['code'] == target['code'] and updated['content']['covers'] == target['content']['covers']


def _endpoint(lead, pid):
    return next(e['id'] for e in lead.get(f'/api/projects/{pid}/endpoints').json()['endpoints'] if e['operationId'] == 'createPet')


def test_regeneration_is_rejected_if_the_case_was_edited_meanwhile(project):
    lead, pid, helper, _, cases = unit_project(project)
    target = cases[0]
    lead.post(f'/api/projects/{pid}/cases/{target["id"]}/regenerate', json={})
    edit = {k: target['content'][k] for k in ('objective', 'preconditions', 'request', 'expected', 'testData', 'cleanup')}
    lead.patch(f'/api/projects/{pid}/unit/cases/{target["id"]}', json={**edit, 'revision': target['revision'], 'title': 'Human edit',
                                                                      'category': target['category'], 'priority': target['priority']})
    status = ClaudeStatus(installed=True, loggedIn=True)
    handlers = {'case.regenerate': lambda job_input, report, submit=None: regenerate.generate(job_input, report, submit, agent=Rewriter())}
    Helper(helper, log=lambda _: None, claude_status=lambda: status, handlers=handlers).run_once()
    assert lead.get(f'/api/projects/{pid}/unit/cases/{target["id"]}').json()['title'] == 'Human edit'
    job = lead.get(f'/api/projects/{pid}/jobs?kind=case.regenerate').json()[0]
    assert job['status'] == 'failed' and 'edited this case' in job['error']


def test_regenerate_flow_retries_until_the_data_flow_is_valid(project):
    lead, pid, helper, _ = setup_project(project)
    lead.post(f'/api/projects/{pid}/flows/generate')
    run_flow_helper(helper, FakeFlowClaude())
    flow = lead.get(f'/api/projects/{pid}/flows').json()['flows'][0]
    lead.post(f'/api/projects/{pid}/cases/{flow["id"]}/regenerate', json={'instruction': 'Verify the pet list after deletion'})
    calls = []

    def writer(agent, task, schema, **options):
        calls.append(task)
        steps = [dict(s) for s in task['currentFlows'][0]['steps']]
        if len(calls) == 1:  # first answer uses a value nobody captured
            steps[1] = {**steps[1], 'request': {**steps[1]['request'], 'pathParams': {'petId': '{{steps.S1.unknown}}'}}}
        else:
            steps.append({**steps[-1], 'key': f'S{len(steps) + 1}', 'operationId': 'listPets', 'purpose': 'List pets after deletion',
                          'expected': {**steps[-1]['expected'], 'status': 200}, 'captures': []})
        return {'flows': [{**task['currentFlows'][0], 'steps': steps, 'title': 'Regenerated flow'}]}

    status = ClaudeStatus(installed=True, loggedIn=True)
    handlers = {'case.regenerate': lambda job_input, report, submit=None: regenerate.generate(job_input, report, submit, agent=writer)}
    assert Helper(helper, log=lambda _: None, claude_status=lambda: status, handlers=handlers).run_once()
    assert len(calls) == 2 and calls[0]['allowStepChanges'] is True
    assert any('before any earlier step captures it' in f['issue'] for f in calls[1]['reviewFindings'])
    updated = next(f for f in lead.get(f'/api/projects/{pid}/flows').json()['flows'] if f['id'] == flow['id'])
    assert updated['title'] == 'Regenerated flow' and updated['content']['operations'][-1] == 'listPets'
    assert len(updated['content']['steps']) == len(flow['content']['steps']) + 1


def test_regenerate_flow_gives_up_after_two_invalid_answers():
    job = {'module': 'flow', 'caseId': 'c1', 'operations': [{'operationId': 'createPet', 'responses': {'201': {}}}],
           'case': {'code': 'TC-0001', 'scenarioKey': 'F01', 'title': 'x', 'category': 'happy_path', 'priority': 'P1',
                    'needsConfirmation': False,
                    'content': {'objective': '', 'preconditions': [], 'testData': '', 'cleanup': [], 'steps': [
                        {'key': 'S1', 'operationId': 'createPet', 'purpose': 'p', 'request': {}, 'expected': {'status': 201}, 'captures': []}]}}}

    def bad(agent, task, schema, **options):
        flow = task['currentFlows'][0]
        return {'flows': [{**flow, 'steps': [{**flow['steps'][0], 'operationId': 'nope'}]}]}
    with pytest.raises(ClaudeError, match='still invalid'):
        regenerate.generate(job, lambda *a: None, agent=bad)
