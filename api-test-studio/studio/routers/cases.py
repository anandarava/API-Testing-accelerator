"""Cross-module case actions: export for sign-off and single-case regeneration."""
import csv
import io
import json
import re
import time
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select
from sqlalchemy.orm import Session as DbSession

from contracts.cases import RegenerateResult, validate_detailed_flow
from contracts.flows import compact_endpoint, outline_from_steps

from ..db import get_db
from ..deps import current_user, membership, record
from ..models import Endpoint, FlowAnalysis, Job, Project, ScriptRecord, TestCase, User
from .helper import helper_state, job_json
from .projects import current_spec
from .unit import case_fingerprint, code, names, project_case, regenerating_ids

router = APIRouter(prefix='/api/projects/{project_id}', tags=['Cases'])
KIND = 'case.regenerate'
STATUS_LABEL = {'in_review': 'In review', 'approved': 'Approved', 'needs_changes': 'Needs changes', 'rejected': 'Rejected'}
PRIORITY_LABEL = {'P1': 'High', 'P2': 'Medium', 'P3': 'Low'}


class RegenerateInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    instruction: str = Field(default='', max_length=2000)


# --- Export -----------------------------------------------------------------------------------------

def _json(value) -> str:
    return '' if value in (None, {}, []) else json.dumps(value, ensure_ascii=False)


def _lines(items) -> str:
    return '\n'.join(f'{i}. {text}' for i, text in enumerate(items or [], 1))


def _checks(expected: dict) -> str:
    return '\n'.join(f"{c['path']} {c['check'].replace('_', ' ')} {_json(c.get('value'))}".strip() for c in expected.get('bodyChecks', []))


def _script_state(case: TestCase, scripted: dict) -> str:
    if case.id not in scripted:
        return 'Not scripted'
    return 'Scripted' if scripted[case.id] == case_fingerprint(case) else 'Script needs update'


def unit_rows(cases, people, scripted):
    header = ['Case ID', 'Endpoint', 'Title', 'Category', 'Priority', 'Status', 'Objective', 'Preconditions', 'Credentials',
              'Path parameters', 'Query', 'Headers', 'Body', 'Expected status', 'Body checks', 'Test data', 'Cleanup',
              'Needs confirmation', 'Confirmation note', 'Review comment', 'Reviewed by', 'Script']
    rows = []
    for c in cases:
        content, request, expected = c.content, c.content.get('request', {}), c.content.get('expected', {})
        rows.append([code(c), f'{c.method} {c.path}', c.title, c.category.replace('_', ' '), PRIORITY_LABEL.get(c.priority, c.priority),
                     STATUS_LABEL[c.status], content.get('objective', ''), _lines(content.get('preconditions')), request.get('auth', ''),
                     _json(request.get('pathParams')), _json(request.get('query')), _json(request.get('headers')),
                     _json(request.get('body')), expected.get('status') or '', _checks(expected), content.get('testData', ''),
                     _lines(content.get('cleanup')), 'Yes' if c.needs_confirmation else 'No', content.get('confirmationNote', ''),
                     c.review_comment, people.get(c.reviewed_by, ''), _script_state(c, scripted)])
    return header, rows


ROLE_KIND = {'allowed': 'Allowed', 'denied': 'Denied', 'ownership': 'Ownership (second user)',
             'hidden_data': 'Hidden data', 'protected_field': 'Protected field'}


def role_rows(cases, people, scripted):
    header = ['Case ID', 'Endpoint', 'Title', 'Check', 'Roles', 'User', 'Priority', 'Status', 'Objective', 'Preconditions',
              'Credentials', 'Path parameters', 'Query', 'Body', 'Expected status', 'Body checks', 'Verify unchanged', 'Cleanup',
              'Needs confirmation', 'Confirmation note', 'Review comment', 'Reviewed by', 'Script']
    rows = []
    for c in cases:
        content, request, expected = c.content, c.content.get('request', {}), c.content.get('expected', {})
        rows.append([code(c), f'{c.method} {c.path}', c.title, ROLE_KIND.get(c.category, c.category), ', '.join(content.get('roles', [])),
                     'Second user' if content.get('userSlot') == 'second' else 'First user',
                     PRIORITY_LABEL.get(c.priority, c.priority), STATUS_LABEL[c.status], content.get('objective', ''),
                     _lines(content.get('preconditions')), request.get('auth', ''), _json(request.get('pathParams')),
                     _json(request.get('query')), _json(request.get('body')), expected.get('status') or '', _checks(expected),
                     content.get('verifyUnchanged', ''), _lines(content.get('cleanup')), 'Yes' if c.needs_confirmation else 'No',
                     content.get('confirmationNote', ''), c.review_comment, people.get(c.reviewed_by, ''), _script_state(c, scripted)])
    return header, rows


def flow_rows(cases, people, scripted, operations):
    header = ['Case ID', 'Title', 'Type', 'Priority', 'Status', 'Goal', 'Preconditions', 'Steps', 'Step count', 'Cleanup',
              'Needs confirmation', 'Confirmation note', 'Review comment', 'Reviewed by', 'Script']
    rows = []
    for c in cases:
        content = c.content
        steps = []
        for i, step in enumerate(content.get('steps', []), 1):
            op = operations.get(step['operationId'], {})
            captures = ', '.join(f"{cap['name']}={cap['path']}" for cap in step.get('captures', []))
            steps.append(f"{i}. {op.get('method', '?')} {op.get('path', step['operationId'])} — {step['purpose']} → "
                         f"{step['expected'].get('status') or '?'}" + (f' (saves {captures})' if captures else ''))
        rows.append([code(c), c.title, c.category.replace('_', ' '), PRIORITY_LABEL.get(c.priority, c.priority), STATUS_LABEL[c.status],
                     content.get('objective') or content.get('goal', ''), _lines(content.get('preconditions')), '\n'.join(steps),
                     len(steps), _lines(content.get('cleanup')), 'Yes' if c.needs_confirmation else 'No',
                     content.get('confirmationNote', ''), c.review_comment, people.get(c.reviewed_by, ''), _script_state(c, scripted)])
    return header, rows


def _xlsx(sheets: list[tuple[str, list, list]], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    for name, header, rows in sheets:
        sheet = book.create_sheet(name)
        sheet.append(header)
        for row in rows:
            sheet.append([ILLEGAL.sub('', v) if isinstance(v, str) else v for v in row])
        for cell in sheet[1]:
            cell.font, cell.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='4F46E5')
        sheet.freeze_panes = 'B2'
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(header, 1):
            longest = max([len(str(column))] + [min(60, max(len(line) for line in str(r[index - 1]).split('\n')) if r[index - 1] != '' else 0) for r in rows])
            sheet.column_dimensions[get_column_letter(index)].width = max(10, min(60, longest + 2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
    book.properties.title = title
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _csv(header, rows) -> bytes:
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(header)
    for row in rows:
        # Neutralise spreadsheet formulas in exported text.
        writer.writerow(["'" + v if isinstance(v, str) and v[:1] in ('=', '+', '-', '@') else v for v in row])
    return ('﻿' + stream.getvalue()).encode('utf-8')


@router.get('/export')
def export(project_id: str, module: Literal['unit', 'flow', 'role', 'all'] = 'all', format: Literal['xlsx', 'csv'] = 'xlsx',
           status: Literal['all', 'in_review', 'approved', 'needs_changes', 'rejected'] = 'all',
           db: DbSession = Depends(get_db), user: User = Depends(current_user)):
    project, _ = membership(db, project_id, user)
    if format == 'csv' and module == 'all':
        raise HTTPException(400, 'CSV exports one module at a time. Choose unit, flow or role.')
    query = select(TestCase).where(TestCase.project_id == project_id)
    if status != 'all':
        query = query.where(TestCase.status == status)
    cases = db.scalars(query.order_by(TestCase.number)).all()
    people = names(db, *[c.reviewed_by for c in cases])
    scripted = {cid: fp for s in db.scalars(select(ScriptRecord).where(ScriptRecord.project_id == project_id))
                for cid, fp in s.case_fingerprints.items()}
    spec = current_spec(db, project_id)
    operations = {e.operation_id: {'method': e.method, 'path': e.path}
                  for e in db.scalars(select(Endpoint).where(Endpoint.spec_id == spec.id))} if spec else {}
    sheets = []
    if module in ('unit', 'all'):
        sheets.append(('Unit cases', *unit_rows([c for c in cases if c.module == 'unit'], people, scripted)))
    if module in ('flow', 'all'):
        sheets.append(('Flow cases', *flow_rows([c for c in cases if c.module == 'flow'], people, scripted, operations)))
    roles = [c for c in cases if c.module == 'role']
    if module == 'role' or (module == 'all' and roles):
        sheets.append(('Role cases', *role_rows(roles, people, scripted)))
    slug = re.sub(r'[^A-Za-z0-9]+', '-', project.name).strip('-').lower() or 'project'
    stamp = time.strftime('%Y-%m-%d')
    scope = module if status == 'all' else f'{module}-{status.replace("_", "-")}'
    filename = f'{slug}-{scope}-cases-{stamp}.{format}'
    body = _xlsx(sheets, f'{project.name} test cases') if format == 'xlsx' else _csv(sheets[0][1], sheets[0][2])
    media = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if format == 'xlsx' else 'text/csv; charset=utf-8'
    total = sum(len(rows) for _, _, rows in sheets)
    record(db, project, user, 'cases.exported', f'{user.name} exported {total} test case{"s" if total != 1 else ""} ({format.upper()})')
    db.commit()
    return Response(body, media_type=media, headers={'Content-Disposition': f'attachment; filename="{filename}"'})


# --- Regeneration -------------------------------------------------------------------------------------

@router.post('/cases/{case_id}/regenerate')
def regenerate(project_id: str, case_id: str, body: RegenerateInput, db: DbSession = Depends(get_db),
               user: User = Depends(current_user)):
    project, _ = membership(db, project_id, user)
    case = project_case(db, project_id, case_id)
    if not helper_state(db, user.id)['connected']:
        raise HTTPException(409, 'Your Local Helper is not connected. Start it on your computer and try again.')
    if case_id in regenerating_ids(db, project_id):
        raise HTTPException(409, f'{code(case)} is already being regenerated.')
    spec = current_spec(db, project_id)
    if not spec:
        raise HTTPException(400, 'Upload an API spec first.')
    endpoints = {e.operation_id: e for e in db.scalars(select(Endpoint).where(Endpoint.spec_id == spec.id))}
    payload = {'caseId': case.id, 'module': case.module, 'instruction': body.instruction.strip(),
               'reviewComment': case.review_comment if case.status == 'needs_changes' else '', 'findings': case.findings,
               'api': {'title': spec.title, 'servers': spec.servers, 'securitySchemes': spec.security_schemes},
               'case': {'code': code(case), 'scenarioKey': case.scenario_key, 'title': case.title, 'category': case.category,
                        'priority': case.priority, 'needsConfirmation': case.needs_confirmation, 'content': case.content,
                        'revision': case.revision}}
    if case.module == 'unit':
        endpoint = next((e for e in endpoints.values() if e.method == case.method and e.path == case.path), None)
        if not endpoint:
            raise HTTPException(400, 'This endpoint is no longer in the current API spec.')
        payload['endpoint'] = endpoint.details
    else:
        used = list(dict.fromkeys(s['operationId'] for s in case.content.get('steps', [])))
        payload['operations'] = [compact_endpoint(e.details) for e in endpoints.values()]
        analysis = db.get(FlowAnalysis, project_id)
        payload['domain'] = analysis.domain if analysis else None
        missing = [op for op in used if op not in endpoints]
        if missing:
            payload['instruction'] += f' The operations {", ".join(missing)} no longer exist; replace them.'
    job = Job(project_id=project_id, kind=KIND, requested_by=user.id, input=payload,
              progress={'message': 'Waiting for your Local Helper', 'percent': 0})
    db.add(job)
    record(db, project, user, 'cases.regeneration_started', f'{user.name} asked Claude to regenerate {code(case)} {case.title}')
    db.commit()
    return job_json(job, user.name)


def store_regenerated(db: DbSession, job: Job, user: User, result: dict):
    parsed = RegenerateResult.model_validate(result)
    case = db.get(TestCase, parsed.caseId)
    if not case or case.id != job.input['caseId']:
        raise HTTPException(400, 'The result is for a different test case.')
    if case.revision != job.input['case']['revision']:
        raise HTTPException(409, 'Someone edited this case while Claude was rewriting it. Regenerate again.')
    content = dict(case.content)
    if case.module == 'unit':
        draft = parsed.unitCase.model_dump() if parsed.unitCase else None
        if not draft:
            raise HTTPException(400, 'A unit case result is required.')
        content.update({k: draft[k] for k in ('objective', 'preconditions', 'request', 'expected', 'testData', 'cleanup')})
        case.category = draft['category']
    else:
        draft = parsed.flowCase.model_dump() if parsed.flowCase else None
        if not draft:
            raise HTTPException(400, 'A flow result is required.')
        spec = current_spec(db, job.project_id)
        operations = {e.operation_id: e.details for e in db.scalars(select(Endpoint).where(Endpoint.spec_id == spec.id))}
        problems = validate_detailed_flow(draft, operations)
        if problems:
            raise HTTPException(400, 'The regenerated flow is invalid: ' + ' '.join(problems[:5]))
        content.update({k: draft[k] for k in ('objective', 'preconditions', 'steps', 'testData', 'cleanup')})
        content['operations'] = [s['operationId'] for s in draft['steps']]
        content['outline'] = outline_from_steps(content.get('outline'), case.scenario_key, draft['title'], draft['type'],
                                                draft['priority'], draft['steps'])
        case.category = draft['type']
    content['confirmationNote'] = draft['confirmationNote']
    case.content, case.title, case.priority = content, draft['title'], draft['priority']
    case.needs_confirmation, case.findings, case.review_comment = draft['needsConfirmation'], [], ''
    case.status, case.revision, case.updated_by, case.updated_at = 'in_review', case.revision + 1, user.id, time.time()
    record(db, db.get(Project, job.project_id), user, 'cases.regenerated',
           f'Claude rewrote {code(case)} {case.title} on {user.name}\'s computer — ready for review')


def complete_regeneration(db: DbSession, job: Job, user: User, result: dict):
    job.result = {'caseId': result.get('caseId')}

