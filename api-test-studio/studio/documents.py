"""Read an uploaded roles file as plain text for Claude, and refuse files that contain credentials."""
import csv
import io
import json
import re
import zipfile
from xml.etree import ElementTree

import yaml

MAX_BYTES = 2 * 1024 * 1024
MAX_CHARS = 80_000
FORMATS = ('.xlsx', '.csv', '.json', '.yaml', '.yml', '.md', '.txt', '.docx')
SECRET_NAME = re.compile(r'^(password|passwd|pwd|pass|secret|client[ _-]?secret|token|access[ _-]?token|api[ _-]?key|apikey|private[ _-]?key)$', re.I)
SECRET_ENV_KEY = re.compile(r'^[A-Za-z0-9_.-]*(PASSWORD|PASSWD|PWD|SECRET|TOKEN|API_?KEY)$', re.I)


class DocumentError(ValueError):
    """A readable reason the file cannot be used."""


CREDENTIALS = ('This file contains credentials ({where}). Keep usernames and passwords in your framework\'s credentials '
               'file (.env) and upload only the roles and what each role may do.')


def _check_rows(rows: list[list[str]], where: str):
    if not rows:
        return
    header = [str(c or '').strip() for c in rows[0]]
    for index, name in enumerate(header):
        if SECRET_NAME.match(name) and any(len(r) > index and str(r[index] or '').strip() for r in rows[1:]):
            raise DocumentError(CREDENTIALS.format(where=f'column "{name}" in {where}'))


def _check_text(text: str):
    for line in text.splitlines():
        match = re.match(r'^\s*(?:export\s+)?([A-Za-z0-9_.-]+)\s*[:=]\s*(\S+)', line)
        if not match:
            continue
        key, value = match.groups()
        if (SECRET_NAME.match(key) or SECRET_ENV_KEY.match(key)) and value.strip('"\'') not in ('', '<', '-', '...', '***'):
            raise DocumentError(CREDENTIALS.format(where=f'"{key}"'))


def _check_data(value, path='file'):
    if isinstance(value, dict):
        for key, item in value.items():
            if (SECRET_NAME.match(str(key)) or SECRET_ENV_KEY.match(str(key))) and isinstance(item, (str, int)) and str(item).strip():
                raise DocumentError(CREDENTIALS.format(where=f'"{key}"'))
            _check_data(item, f'{path}.{key}')
    elif isinstance(value, list):
        for item in value:
            _check_data(item, path)


def _table(rows: list[list]) -> str:
    return '\n'.join(' | '.join('' if c is None else str(c).strip() for c in row) for row in rows if any(c not in (None, '') for c in row))


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for damaged files
        raise DocumentError('The Excel file could not be read. Save it again as .xlsx and retry.') from exc
    parts = []
    for sheet in book.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        _check_rows([[str(c or '') for c in r] for r in rows], f'sheet "{sheet.title}"')
        if rows:
            parts.append(f'## Sheet: {sheet.title}\n{_table(rows)}')
    return '\n\n'.join(parts)


def _docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            root = ElementTree.fromstring(archive.read('word/document.xml'))
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError) as exc:
        raise DocumentError('The Word file could not be read. Save it again as .docx and retry.') from exc
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    lines = []
    body = root.find('w:body', ns)
    for block in list(body) if body is not None else []:
        if block.tag.endswith('}tbl'):
            rows = [[''.join(t.text or '' for t in cell.iter(f"{{{ns['w']}}}t")) for cell in row.findall('w:tc', ns)]
                    for row in block.findall('w:tr', ns)]
            _check_rows(rows, 'a table')
            lines.append(_table(rows))
        else:
            text = ''.join(t.text or '' for t in block.iter(f"{{{ns['w']}}}t"))
            if text.strip():
                lines.append(text)
    return '\n'.join(lines)


def extract_text(filename: str, data: bytes) -> str:
    """The roles file as text Claude can read. Raises DocumentError with a reason the user can act on."""
    name = filename.lower()
    if not name.endswith(FORMATS):
        raise DocumentError(f'Upload the roles file as {", ".join(FORMATS)}.')
    if len(data) > MAX_BYTES:
        raise DocumentError('The roles file is larger than 2 MB. Keep only the roles and their permissions.')
    if not data.strip():
        raise DocumentError('The roles file is empty.')
    if name.endswith('.xlsx'):
        text = _xlsx(data)
    elif name.endswith('.docx'):
        text = _docx(data)
    else:
        try:
            raw = data.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise DocumentError('The file is not UTF-8 text. Save it as UTF-8 and retry.') from exc
        if name.endswith('.csv'):
            rows = list(csv.reader(io.StringIO(raw)))
            _check_rows(rows, 'the CSV header')
            text = _table(rows)
        elif name.endswith(('.json', '.yaml', '.yml')):
            try:
                parsed = json.loads(raw) if name.endswith('.json') else yaml.safe_load(raw)
            except (ValueError, yaml.YAMLError) as exc:
                raise DocumentError(f'The file is not valid {"JSON" if name.endswith(".json") else "YAML"}: {str(exc)[:200]}') from exc
            _check_data(parsed)
            text = yaml.safe_dump(parsed, sort_keys=False, allow_unicode=True)
        else:
            _check_text(raw)
            text = raw
    text = text.strip()
    if not text:
        raise DocumentError('No text was found in the roles file.')
    if len(text) > MAX_CHARS:
        raise DocumentError(f'The roles file has {len(text):,} characters; keep it under {MAX_CHARS:,} '
                            '(the roles and what each may do are enough).')
    return text
